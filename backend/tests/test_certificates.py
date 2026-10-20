import io
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import app
from auth import require_user
from diasavto_validator.certificates import CertificateFormatError, extract_masks
from diasavto_validator.hi_lex import validate_catalog
from diasavto_validator.models import CatalogItem


def test_reads_txt_masks() -> None:
    assert extract_masks(b"4VB**** 9BB**** 4VB****", "mask.txt") == (
        "4VB****",
        "9BB****",
    )


def test_rejects_unknown_mask() -> None:
    with pytest.raises(CertificateFormatError, match="неутвержд"):
        extract_masks(b"4QB****", "mask.txt")


def test_rejects_missing_masks() -> None:
    with pytest.raises(CertificateFormatError, match="не найдены"):
        extract_masks(b"no masks", "mask.txt")


def test_extracts_docx_masks() -> None:
    buffer = io.BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body><w:p><w:r><w:t>4VB**** 9BB****</w:t></w:r></w:p></w:body>"
            "</w:document>",
        )
    assert extract_masks(buffer.getvalue(), "mask.docx") == ("4VB****", "9BB****")


def test_filters_by_certificate_masks() -> None:
    items = [
        CatalogItem("4VB0000", "HI-LEX CORPORATION", "тросы стояночного тормоза"),
        CatalogItem("9BB0000", "HI-LEX CORPORATION", "тросы стояночного тормоза"),
    ]
    matched = validate_catalog(items, ("4VB****",))
    assert matched[0].is_accepted
    assert not matched[1].is_accepted


def test_api_applies_certificate_masks() -> None:
    app.dependency_overrides[require_user] = lambda: "test-user"
    try:
        catalog = (
            "article,manufacturer,product_group\n"
            "4VB0000,HI-LEX CORPORATION,тросы стояночного тормоза\n"
            "9BB0000,HI-LEX CORPORATION,тросы стояночного тормоза\n"
        ).encode("utf-8")
        result = TestClient(app).post(
            "/api/validate",
            files=[
                ("file", ("catalog.csv", catalog, "text/csv")),
                ("certificate", ("mask.txt", b"4VB****", "text/plain")),
            ],
        )
        assert result.status_code == 200
        book = load_workbook(io.BytesIO(result.content), read_only=True)
        assert book.worksheets[0].max_row == 2
        assert book.worksheets[1].max_row == 2
    finally:
        app.dependency_overrides.clear()
