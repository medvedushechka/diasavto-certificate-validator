import csv
import io
from collections.abc import Iterable
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import CatalogItem, ValidationResult

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_ROWS = 10_000
REQUIRED_COLUMNS = ("article", "manufacturer", "product_group")


class CatalogFormatError(ValueError):
    pass


def _validate_headers(headers: Iterable[object]) -> dict[str, int]:
    names = [value.strip().lower() if isinstance(value, str) else "" for value in headers]
    if any(names.count(name) != 1 for name in REQUIRED_COLUMNS):
        raise CatalogFormatError("Нужны уникальные столбцы: article, manufacturer, product_group")
    return {name: names.index(name) for name in REQUIRED_COLUMNS}


def _parse_rows(rows: Iterable[tuple[object, ...]], headers: dict[str, int]) -> list[CatalogItem]:
    items: list[CatalogItem] = []
    for number, row in enumerate(rows, start=2):
        if not any(value is not None and str(value) != "" for value in row):
            continue
        if number > MAX_ROWS + 1:
            raise CatalogFormatError(f"Максимум {MAX_ROWS} строк каталога")
        values: list[str] = []
        for name in REQUIRED_COLUMNS:
            index = headers[name]
            if index >= len(row) or row[index] is None or not isinstance(row[index], str):
                raise CatalogFormatError(f"Строка {number}: поле {name} должно быть текстом")
            values.append(row[index])
        items.append(CatalogItem(*values))
    if not items:
        raise CatalogFormatError("Каталог не содержит товаров")
    return items


def parse_catalog(content: bytes, filename: str) -> list[CatalogItem]:
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise CatalogFormatError("Файл пустой или превышает 2 МБ")
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        try:
            reader = csv.reader(io.StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
            headers = _validate_headers(next(reader))
            return _parse_rows(reader, headers)
        except (UnicodeError, csv.Error, StopIteration) as exc:
            raise CatalogFormatError("Некорректный CSV. Используйте UTF-8") from exc
    if extension == ".xlsx":
        if not content.startswith(b"PK\x03\x04"):
            raise CatalogFormatError("Некорректный XLSX")
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                if sheet is None:
                    raise CatalogFormatError("В книге нет активного листа")
                if sheet.max_row is not None and sheet.max_row > MAX_ROWS + 1:
                    raise CatalogFormatError(f"Максимум {MAX_ROWS} строк каталога")
                rows = iter(sheet.values)
                headers = _validate_headers(next(rows))
                return _parse_rows(rows, headers)
            finally:
                workbook.close()
        except (
            InvalidFileException,
            OSError,
            ValueError,
            KeyError,
            StopIteration,
            EOFError,
        ) as exc:
            if isinstance(exc, CatalogFormatError):
                raise
            raise CatalogFormatError("Повреждённый или пустой XLSX") from exc
    raise CatalogFormatError("Допускаются только CSV и XLSX")


def _safe_excel_cell(value: str) -> str:
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def create_report(results: list[ValidationResult]) -> bytes:
    workbook = Workbook()
    accepted = workbook.active
    if accepted is None:
        raise RuntimeError("Не удалось создать лист отчёта")
    accepted.title = "Подходят"
    rejected = workbook.create_sheet("Исключения")
    headers = ("article", "manufacturer", "product_group", "search_key", "mask", "reason")
    accepted.append(headers)
    rejected.append(headers)
    for result in results:
        item = result.item
        row = (
            _safe_excel_cell(item.article),
            _safe_excel_cell(item.manufacturer),
            _safe_excel_cell(item.product_group),
            _safe_excel_cell(result.search_key),
            result.matched_mask or "",
            result.rejection_reason.value if result.rejection_reason else "",
        )
        (accepted if result.is_accepted else rejected).append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
